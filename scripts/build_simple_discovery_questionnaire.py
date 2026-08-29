from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_discovery_questionnaire import (
    INK,
    LINE,
    MUTED,
    NAVY,
    PALE_BLUE,
    RED,
    WHITE,
    add_page_number,
    add_question,
    add_section_heading,
    add_simple_table,
    keep_paragraph,
    set_cell_border,
    set_cell_margins,
    set_cell_shading,
    set_run_font,
)


OUTPUT = Path(
    r"C:\Users\divin\Downloads\BestWorth\output\documents\BestWorth_Simple_Discovery_Questionnaire.docx"
)


def populate_header(header):
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    table.autofit = False
    table.columns[0].width = Inches(4.7)
    table.columns[1].width = Inches(1.8)
    left, right = table.rows[0].cells
    left.text = "BESTWORTH PRODUCTS LIMITED"
    right.text = "PROJECT QUESTIONNAIRE"
    for run in left.paragraphs[0].runs:
        set_run_font(run, 8, True, NAVY)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in right.paragraphs[0].runs:
        set_run_font(run, 7.5, True, RED)
    for cell in (left, right):
        set_cell_border(cell, bottom={"val": "single", "sz": 8, "color": LINE})
        set_cell_margins(cell, top=20, bottom=60, start=0, end=0)


def populate_footer(footer):
    paragraph = footer.paragraphs[0]
    paragraph.text = "Confidential • BestWorth project discovery • 24 August 2026"
    for run in paragraph.runs:
        set_run_font(run, 7.5, False, MUTED)
    add_page_number(footer.add_paragraph())


def build_document():
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.color.rgb = RGBColor.from_string(NAVY)

    populate_header(section.header)
    populate_header(section.even_page_header)
    populate_footer(section.footer)
    populate_footer(section.even_page_footer)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run("PROJECT DISCOVERY")
    set_run_font(run, 8, True, RED)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run("Builders Mart & Product Breakdown Questionnaire")
    set_run_font(run, 19, True, NAVY)
    keep_paragraph(paragraph, True)

    paragraph = doc.add_paragraph(
        "A short form to confirm what the client wants before the website section and admin tool are designed."
    )
    paragraph.paragraph_format.space_after = Pt(12)
    for run in paragraph.runs:
        set_run_font(run, 10, False, MUTED)

    meta = doc.add_table(rows=2, cols=4)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    fields = [
        ("Completed by", "Type name"),
        ("Role", "Type role"),
        ("Date", "DD / MM / YYYY"),
        ("Contact", "Type email or phone"),
    ]
    for index, (label, value) in enumerate(fields):
        row = index // 2
        column = (index % 2) * 2
        label_cell = meta.rows[row].cells[column]
        value_cell = meta.rows[row].cells[column + 1]
        label_cell.width = Inches(1.15)
        value_cell.width = Inches(2.1)
        set_cell_shading(label_cell, PALE_BLUE)
        set_cell_shading(value_cell, WHITE)
        for cell in (label_cell, value_cell):
            set_cell_margins(cell)
            set_cell_border(cell, bottom={"val": "single", "sz": 5, "color": LINE})
        label_cell.text = label.upper()
        value_cell.text = value
        for run in label_cell.paragraphs[0].runs:
            set_run_font(run, 7.5, True, NAVY)
        for run in value_cell.paragraphs[0].runs:
            run.italic = True
            set_run_font(run, 8.5, False, MUTED)

    doc.add_paragraph()
    note = doc.add_table(rows=1, cols=1)
    cell = note.cell(0, 0)
    set_cell_shading(cell, PALE_BLUE)
    set_cell_margins(cell, top=125, start=160, bottom=125, end=160)
    set_cell_border(cell, start={"val": "single", "sz": 20, "color": RED})
    cell.text = "Tick the closest answer and add a short explanation where needed. Write N/A for anything that does not apply. Do not include complete bank account numbers in this form."
    for run in cell.paragraphs[0].runs:
        set_run_font(run, 9, False, INK)

    add_section_heading(doc, "01", "Company & Builders Mart", "Confirm the public identity and purpose of the new retail section.")
    add_question(doc, "1.1", "What is the exact company name that should appear on the website and official documents?", lines=2, critical=True)
    add_question(doc, "1.2", "Is “BestWorth Builders Mart” the confirmed name?", ["Yes", "No", "Name still being decided"], lines=1, critical=True)
    add_question(doc, "1.3", "How is Builders Mart connected to BestWorth Products Limited?", ["Retail department", "Division", "Separate company", "Other"], lines=1)
    add_question(doc, "1.4", "What should the Builders Mart section achieve?", ["Introduce the retail outlet", "Show available products", "Receive enquiries", "Help customers contact or visit the store"], lines=1, critical=True)
    add_question(doc, "1.5", "Provide the shop address, phone/WhatsApp, opening hours and service area.", lines=4)

    add_section_heading(doc, "02", "Website section", "Confirm where it belongs and what visitors should see.")
    add_question(doc, "2.1", "Where should Builders Mart appear on the homepage?", ["Before News & Media", "After News & Media", "Before Products", "Other"], lines=1, critical=True)
    add_question(doc, "2.2", "In the voice note, what did “before the placeholder” mean? Name the exact section.", lines=2, critical=True)
    add_question(doc, "2.3", "Should Builders Mart have a separate full page?", ["Homepage section only", "Homepage section and full page", "Full page only"], lines=1)
    add_question(doc, "2.4", "What should visitors see in the section?", ["Short introduction", "Store photos", "Selected products", "Location/map", "Contact button"], lines=1)
    add_question(doc, "2.5", "Should product prices be shown publicly?", ["No", "Yes", "Only selected prices", "Not decided"], lines=1)

    add_section_heading(doc, "03", "Products & sizes", "List the products and details staff need when preparing a breakdown.")
    add_question(doc, "3.1", "Which products should be available first?", ["Nails", "Cement", "Wire mesh", "Screws", "Other"], lines=1)
    add_question(doc, "3.2", "Confirm the nail sizes.", ["1 inch", "1.5 inches", "2 inches", "2.5 inches", "3 inches", "4 inches", "5 inches", "6 inches"], lines=1)
    add_question(doc, "3.3", "For each product, what unit is sold?", lines=3, note="Examples: bag, piece, carton, roll, bundle or tonne.")
    add_question(doc, "3.4", "Complete the starter product list below or attach the approved list.", lines=0)
    add_simple_table(
        doc,
        ["Product", "Available sizes / types", "Unit sold", "Notes"],
        [["Type product", "List sizes or types", "Type unit", "Type note"] for _ in range(6)],
        [1.35, 2.15, 1.2, 1.8],
    )

    add_section_heading(doc, "04", "Product breakdown document", "Confirm the simple workflow used to prepare and send the customer document.")
    add_question(doc, "4.1", "What should the document be called?", ["Product Breakdown", "Quotation", "Proforma Invoice", "Other"], lines=1, critical=True)
    add_question(doc, "4.2", "Who will prepare it?", ["Owner", "Sales workers", "Accounts workers", "Selected workers"], lines=1)
    add_question(doc, "4.3", "Which customer details are required?", ["Name", "Company", "Email", "Phone", "Address"], lines=1)
    add_question(doc, "4.4", "Should the reference/inquiry number be automatic or entered by staff?", ["Automatic", "Entered by staff", "Either"], lines=1, critical=True)
    add_question(doc, "4.5", "How should prices work?", ["Staff enter every price", "Use saved prices", "Saved prices that staff can change"], lines=1, critical=True)
    add_question(doc, "4.6", "Which amounts should the document calculate?", ["Quantity × price", "Line totals", "Delivery", "Discount", "VAT/tax", "Final total"], lines=1)
    add_question(doc, "4.7", "How should the completed document be sent?", ["PDF by email", "Download PDF", "Print", "WhatsApp attachment"], lines=1, critical=True)
    add_question(doc, "4.8", "Attach an existing quotation, breakdown or letterhead that shows the preferred design.", ["Attached", "Will send later", "None available"], lines=1)

    add_section_heading(doc, "05", "Payment, approval & staff access", "Keep payment details controlled and confirm who may use the new admin page.")
    add_question(doc, "5.1", "Should approved bank details appear on the document?", ["Always", "Staff chooses when needed", "No"], lines=1, critical=True)
    add_question(doc, "5.2", "Who may manage the saved bank details?", ["Owner only", "Owner and selected accounts worker"], lines=1, critical=True)
    add_question(doc, "5.3", "Should a worker’s breakdown require approval before it is sent?", ["Yes", "No", "Only for selected workers"], lines=1)
    add_question(doc, "5.4", "Who should have access to the Product Breakdown admin page?", ["Owner only", "Selected workers with View access", "Selected workers with Manage access"], lines=1, critical=True)
    add_question(doc, "5.5", "Who gives final approval for the website section and admin tool?", lines=2)

    add_section_heading(doc, "06", "Materials & confirmation", "Confirm what the client will supply before design begins.")
    add_question(doc, "6.1", "Which materials will be supplied?", ["Official logo", "Letterhead", "Product list", "Price list", "Store photos", "Contact details", "Sample quotation"], lines=1)
    add_question(doc, "6.2", "List anything important that is missing from this questionnaire.", lines=4)
    add_question(doc, "6.3", "Can these answers be used as the approved basis for planning?", ["Yes", "Yes, with listed corrections", "No—another discussion is required"], lines=2, critical=True)

    sign = doc.add_table(rows=2, cols=4)
    sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign.autofit = False
    values = [
        ("Approved by", "Type full name", "Role", "Type role"),
        ("Confirmation", "Type name or sign", "Date", "DD / MM / YYYY"),
    ]
    for row, row_values in zip(sign.rows, values):
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            set_cell_margins(cell, top=90, bottom=90)
            set_cell_border(cell, bottom={"val": "single", "sz": 5, "color": LINE})
            cell.text = value
            if index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
                for run in cell.paragraphs[0].runs:
                    set_run_font(run, 8, True, NAVY)
            else:
                set_cell_shading(cell, WHITE)
                for run in cell.paragraphs[0].runs:
                    run.italic = True
                    set_run_font(run, 8.5, False, MUTED)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("END OF QUESTIONNAIRE")
    set_run_font(run, 8, True, RED)

    doc.core_properties.title = "BestWorth Builders Mart and Product Breakdown Questionnaire"
    doc.core_properties.subject = "Simplified project discovery questionnaire"
    doc.core_properties.author = "BestWorth Products Limited"
    doc.core_properties.keywords = "BestWorth, Builders Mart, product breakdown, questionnaire"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
